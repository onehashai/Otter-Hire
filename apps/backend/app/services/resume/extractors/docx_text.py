"""DOCX text: paragraphs + table cells (python-docx), OCR on embedded images if needed."""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
from io import BytesIO
from zipfile import ZipFile

from docx import Document

from app.services.resume.hyperlinks import ResumeHyperlink

from . import ocr


def _normalize_whitespace(text: str) -> str:
    return re.sub(r"\s+", " ", (text or "")).strip()


def _is_low_text(value: str, threshold: int = 80) -> bool:
    return len(_normalize_whitespace(value)) < threshold


def _table_lines(document: Document) -> list[str]:
    lines: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [((c.text or "").strip()) for c in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                lines.append(line)
    return lines


def _extract_docx_hyperlinks(content: bytes) -> list[ResumeHyperlink]:
    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    hyperlinks: list[ResumeHyperlink] = []
    with ZipFile(BytesIO(content)) as zf:
        try:
            rels_xml = zf.read("word/_rels/document.xml.rels")
            doc_xml = zf.read("word/document.xml")
        except KeyError:
            return hyperlinks

    rels_root = ET.fromstring(rels_xml)
    rel_targets = {
        rel.attrib.get("Id"): rel.attrib.get("Target", "")
        for rel in rels_root.findall(".//pr:Relationship", ns)
        if rel.attrib.get("TargetMode") == "External"
    }

    doc_root = ET.fromstring(doc_xml)
    for p_index, paragraph in enumerate(doc_root.findall(".//w:p", ns), start=1):
        for link in paragraph.findall(".//w:hyperlink", ns):
            rel_id = link.attrib.get(f"{{{ns['r']}}}id")
            target = rel_targets.get(rel_id or "")
            if not target:
                continue
            texts = [
                (node.text or "").strip()
                for node in link.findall(".//w:t", ns)
                if (node.text or "").strip()
            ]
            label = " ".join(texts).strip()
            hyperlinks.append(
                ResumeHyperlink(
                    label=label,
                    target=target,
                    source_format="docx",
                    source_hint=f"paragraph:{p_index}",
                )
            )
    return hyperlinks


def extract_docx_plain_text(content: bytes) -> tuple[str, list[str], list[ResumeHyperlink]]:
    warnings: list[str] = []
    document = Document(BytesIO(content))
    hyperlinks = _extract_docx_hyperlinks(content)
    lines: list[str] = []
    for p in document.paragraphs:
        t = (p.text or "").strip()
        if t:
            lines.append(t)
    lines.extend(_table_lines(document))
    parsed = "\n".join(lines).strip()

    if _is_low_text(parsed):
        warnings.append("docx_low_text_trying_ocr_images")
        try:
            ocr_text = ocr.ocr_docx_embedded_images(content)
            if ocr_text:
                warnings.append("ocr_docx_images_used")
                return ocr_text, warnings, hyperlinks
        except Exception:
            warnings.append("ocr_docx_failed")

    return parsed, warnings, hyperlinks
